import json,os
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path.cwd(); out=root/'설명서';out.mkdir(exist_ok=True)
pages=json.loads((root/'설명서/원본/설명서_본문.json').read_text(encoding='utf8'))
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
sec.header_distance=sec.footer_distance=Inches(.3);sec.different_first_page_header_footer=True
styles=doc.styles
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=styles[n];s.font.name='맑은 고딕';s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
 for attr in list(s.element.rPr.rFonts.attrib):
  if attr.endswith('Theme'):del s.element.rPr.rFonts.attrib[attr]
 snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');s.element.get_or_add_pPr().append(snap)
 s.font.color.rgb=RGBColor.from_string('242538')
styles['Normal'].font.size=Pt(12)
styles['Normal'].paragraph_format.line_spacing=Pt(19)
styles['Normal'].paragraph_format.space_after=Pt(8)
styles['Title'].font.size=Pt(34);styles['Title'].font.bold=True;styles['Title'].font.color.rgb=RGBColor(0,0,0)
styles['Heading 1'].font.size=Pt(23);styles['Heading 1'].font.bold=True;styles['Heading 1'].paragraph_format.space_after=Pt(16);styles['Heading 1'].paragraph_format.line_spacing=Pt(32)
styles['Heading 2'].font.size=Pt(14);styles['Heading 2'].font.bold=True;styles['Heading 2'].font.color.rgb=RGBColor.from_string('58417F');styles['Heading 2'].paragraph_format.space_before=Pt(12);styles['Heading 2'].paragraph_format.space_after=Pt(6);styles['Heading 2'].paragraph_format.line_spacing=Pt(21)
styles['Caption'].font.size=Pt(10);styles['Caption'].font.color.rgb=RGBColor.from_string('626478');styles['Caption'].paragraph_format.space_after=Pt(12)
header=sec.header.paragraphs[0];header.text='우주 교실  |  선생님 사용 설명서';header.style='Caption'
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('우주ssam   ·   ');r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('626478')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='우주 교실 선생님 사용 설명서';doc.core_properties.author='우주ssam';doc.core_properties.subject='교실 입장 학급 운영 학습과 놀이 활동 안내';doc.core_properties.keywords='우주 교실 선생님 사용 설명서'
p=doc.add_paragraph('선생님을 위한 기능 안내',style='Subtitle');p.paragraph_format.space_before=Pt(35);p.paragraph_format.space_after=Pt(18)
doc.add_paragraph('우주 교실\n선생님 사용 설명서',style='Title')
p=doc.add_paragraph('교실 입장부터 과제와 미션 운영까지');p.paragraph_format.space_after=Pt(22)
doc.add_picture(str(root/'client/assets/lobby-milky-way.png'),width=Inches(7.0))
doc.paragraphs[-1].paragraph_format.line_spacing=1
p=doc.add_paragraph('2026년 10월 11일 기준\n제작자 우주ssam');p.paragraph_format.space_before=Pt(20);p.paragraph_format.line_spacing=1.5
doc.add_paragraph('수업 준비 · 학급 관리 · 학습 활동 · 놀이와 보상',style='Caption')
def para(text,style=None):
 p=doc.add_paragraph(text,style);p.paragraph_format.widow_control=True;return p
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[2.0,5.06] if len(headers)==2 else [1.4,2.4,3.26]
 if len(headers)==3 and headers[0]=='시간':widths=[.8,2.7,3.56]
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for idx,c in enumerate(t.rows[0].cells):c.text=headers[idx]
 for row in rows:
  for c,text in zip(t.add_row().cells,row):c.text=text
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
   tcpr.append(b);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'263E60' if ri==0 else ('F3F5F9' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
   margins=OxmlElement('w:tcMar')
   for edge,value in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
   tcpr.append(margins)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(16)
    for r in p.runs:r.font.size=Pt(11);r.font.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '242538')
 p=para('');p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=Pt(2);p.paragraph_format.space_before=Pt(0)
for i,page in enumerate(pages,2):
 h=doc.add_heading(page['title'],1);h.paragraph_format.page_break_before=True
 if page.get('toc'):
  para('필요한 기능부터 찾아 읽으세요. 쪽수는 표지를 포함한 책자 기준입니다.')
  for j,p in enumerate(pages,2):
   if j<=3:continue
   par=para(p['title']);par.paragraph_format.space_after=Pt(3);par.paragraph_format.line_spacing=Pt(17)
   tab=par.paragraph_format.tab_stops;from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
   tab.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS);par.add_run('\t'+str(j))
  continue
 if page.get('lead'):para(page['lead']).paragraph_format.space_after=Pt(14)
 if page.get('image'):
  from PIL import Image
  im=Image.open(root/page['image']);w,h=im.size;maxh=2.1 if i==8 else 2.8;wi=min(7.0,maxh*w/h)
  pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.line_spacing=1;pp.add_run().add_picture(str(root/page['image']),width=Inches(wi))
  para(page['caption'],'Caption')
 for b in page.get('blocks',[]):
  if b[0]=='h':doc.add_heading(b[1],2)
  elif b[0]=='p':para(b[1])
  elif b[0]=='table':table(b[1],b[2])
  elif b[0] in ('steps','bullets'):
   for k,line in enumerate(b[1],1):
    p=para((str(k)+'. ' if b[0]=='steps' else '• ')+line);p.paragraph_format.left_indent=Inches(.17);p.paragraph_format.first_line_indent=Inches(-.17);p.paragraph_format.space_after=Pt(5 if b[0]=='steps' else 3)
for tree in [doc.element,doc.styles.element]:
 for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(out/'우주교실_선생님_사용설명서_2026-10-11.docx')
print('Created manual with planned pages:',len(pages)+1)
